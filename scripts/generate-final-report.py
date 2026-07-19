from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "烽燧智守_项目附加报告_最终版.docx"
SHOTS = ROOT / "deliverables" / "screenshots"

NAVY = "0B2545"
BLUE = "1F4D78"
CYAN = "0E7490"
LIGHT = "E8EEF5"
LIGHTER = "F4F6F9"
MUTED = "566573"
RED = "9B1C1C"
GOLD = "7A5A00"
WHITE = "FFFFFF"


def set_east_asia(run, font="Microsoft YaHei"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, LIGHT)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_east_asia(run)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(NAVY)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_east_asia(run)
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string("263746")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, title, body, color=CYAN):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [6.5])
    shade(table.cell(0, 0), LIGHTER)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_east_asia(r)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.size = Pt(10)
    p2 = table.cell(0, 0).add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_east_asia(r2)
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor.from_string("263746")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc, filename, caption):
    path = SHOTS / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Inches(6.28))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    run = cp.add_run(caption)
    set_east_asia(run)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_east_asia(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_east_asia(r)
    return p


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("263746")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_specs = {
        "Heading 1": (16, 16, 8, BLUE),
        "Heading 2": (13, 12, 6, BLUE),
        "Heading 3": (12, 8, 4, NAVY),
    }
    for name, (size, before, after, color) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "烽燧智守  |  AI大赛项目附加报告"
    hp.paragraph_format.space_after = Pt(0)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in hp.runs:
        set_east_asia(run)
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = fp.add_run("比赛演示规则模型  |  第 ")
    set_east_asia(label)
    label.font.size = Pt(8)
    label.font.color.rgb = RGBColor.from_string(MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    tail = fp.add_run(" 页")
    set_east_asia(tail)
    tail.font.size = Pt(8)
    tail.font.color.rgb = RGBColor.from_string(MUTED)


def title_page(doc):
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kr = kicker.add_run("AI先锋未来人才大赛 · 项目附加报告")
    set_east_asia(kr)
    kr.bold = True
    kr.font.size = Pt(11)
    kr.font.color.rgb = RGBColor.from_string(CYAN)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(12)
    tr = title.add_run("烽燧智守")
    set_east_asia(tr)
    tr.bold = True
    tr.font.size = Pt(32)
    tr.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("面向垃圾焚烧发电厂关键设备的故障预警与智能运维系统")
    set_east_asia(sr)
    sr.font.size = Pt(15)
    sr.font.color.rgb = RGBColor.from_string(BLUE)
    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    callout.style = "Table Grid"
    set_table_geometry(callout, [6.5])
    shade(callout.cell(0, 0), LIGHTER)
    p = callout.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("可运行 · 可解释 · 可协同 · 可闭环 · 可逐步接入真实数据")
    set_east_asia(r)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(CYAN)
    for _ in range(5):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("最终成品版  |  2026年7月")
    set_east_asia(mr)
    mr.font.size = Pt(10)
    mr.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_page_break()


def build_report():
    doc = Document()
    configure_document(doc)
    title_page(doc)

    doc.add_heading("执行摘要", level=1)
    doc.add_paragraph("烽燧智守以垃圾焚烧发电厂关键旋转设备为对象，把监测、工况识别、健康评估、异常预警、知识检索、Agent编排、三维定位、飞书工单、备件扣减、维修验证和知识沉淀连接为可追溯业务闭环。比赛版本提供12台设备、10个故障案例、9类语义3D模型和完整Mock演示；真实平台能力按配置逐模块启用。")
    add_callout(doc, "真实性声明", "遥测、健康度、故障概率、诊断结论和经济价值均为模拟演示或目标口径；当前规则模型不代表生产准确率，是否停机必须由现场负责人依据安全规程决定。", GOLD)
    add_table(doc, ["交付维度", "本次成品"], [
        ["业务闭环", "1号引风机68→92、预警关闭、库存6→5、知识候选"],
        ["AI能力", "本地RAG、十工具Agent、结构化Provider、多模态辅助与失败降级"],
        ["飞书工具链", "多维表格工单、机器人/卡片/事件/WebSocket适配，Aily与妙搭指南"],
        ["数字孪生", "增强/原始引风机、三档LOD、八类组件模型、节点与故障联动"],
        ["工程质量", "TypeScript strict、Vitest/Supertest/Playwright、离线演示、Vercel/EdgeOne"],
    ], [1.4, 5.1])
    add_figure(doc, "01-dashboard.png", "图1  设备总览：12台演示设备与垃圾焚烧厂关键设备拓扑")

    doc.add_heading("1. 需求与业务洞察", level=1)
    doc.add_paragraph("运行值班、设备维修、设备管理和工厂管理层关注的不是同一张报表，但共享同一条证据链。系统以风险处置旅程组织功能，避免把AI做成孤立聊天入口。")
    add_table(doc, ["角色", "核心痛点", "系统价值"], [
        ["运行值班", "多系统切换、告警缺少工况语境", "快速理解风险、证据和建议时限"],
        ["设备维修", "资料分散、检查清单不完整", "3D部件定位、RAG步骤和移动巡检入口"],
        ["设备管理", "工单、库存和历史案例割裂", "受控Agent、库存联动和知识候选"],
        ["管理层", "指标口径不一、责任和进度不可见", "状态总览、闭环轨迹和数据边界"],
    ], [1.25, 2.35, 2.9])
    doc.add_heading("替代方案差异", level=2)
    add_table(doc, ["维度", "固定阈值", "普通设备管理", "单一诊断", "烽燧智守"], [
        ["工况适应", "弱", "弱", "视产品", "五类工况基准"],
        ["可解释证据", "阈值", "流程", "视模型", "趋势+贡献+引用"],
        ["3D/RAG/Agent", "无", "通常无", "局部", "统一接入"],
        ["工单与备件", "弱", "强", "弱", "闭环联动"],
        ["部署路径", "较低", "中高", "高", "Mock起步、模块化启用"],
    ], [1.1, 1.2, 1.35, 1.2, 1.65])

    doc.add_page_break()
    doc.add_heading("2. 系统架构与业务闭环", level=1)
    add_callout(doc, "闭环主线", "设备数据 → 工况基准 → 健康与风险 → RAG证据 → Agent工具编排 → 3D定位 → 库存检查 → 飞书工单 → 维修验证 → 知识沉淀")
    add_table(doc, ["层级", "职责", "关键实现"], [
        ["体验层", "深色工业指挥UI、飞书网页应用、普通浏览器", "React、Router、Query、响应式设计"],
        ["智能层", "意图、RAG、Agent、结构化诊断、多模态", "Zod、轻量索引、十工具、降级Provider"],
        ["业务层", "健康评估、预警、工单、库存、知识", "可解释规则、状态机、幂等和审计日志"],
        ["适配层", "Mock/Feishu、通知、认证和模型Provider", "Repository/Provider接口与分模块能力判断"],
        ["数据与资产", "演示数据、多维表格、GLB/Blend和文档", "固定种子、localStorage重放、LOD模型库"],
    ], [1.1, 2.0, 3.4])

    doc.add_heading("3. 六大核心能力", level=1)
    capabilities = [
        ("AI算法基础", "工况规则、RAG、Agent、结构化输出、多模态、失败降级和可解释证据链"),
        ("AI工具链应用", "飞书多维表格、机器人、卡片、WebSocket、Aily、妙搭和豆包Provider"),
        ("工程落地能力", "Monorepo、统一领域模型、API、测试、安全、性能、离线与部署回滚"),
        ("产品思维", "四类用户、风险旅程、人工确认、模块优先级和分阶段落地"),
        ("业务洞察", "垃圾焚烧关键设备、10类故障、多参数贡献、工单/库存/知识联动"),
        ("表达展示", "深色工业UI、拓扑、语义3D、截图、报告、PPT、路演稿和Q&A"),
    ]
    for title, text in capabilities:
        add_bullet(doc, f"{title}：{text}")
    add_figure(doc, "15-six-skills-map.png", "图2  六大能力映射页：功能、技术、Demo与证据文件对应")

    doc.add_page_break()
    doc.add_heading("4. AI算法与知识能力", level=1)
    doc.add_heading("可解释健康评估", level=2)
    doc.add_paragraph("系统按停机、启动、低负荷、稳定、高负荷切换基准，对振动、温度、电流、压力和转速计算偏离并归一化，再按设备类型赋权。健康度与关键指标越限共同决定风险，下降贡献可以回溯到指标证据。")
    doc.add_heading("可追溯RAG", level=2)
    add_number(doc, "将知识条目、10个故障案例和历史工单转为带来源、设备、故障和风险元数据的文档。")
    add_number(doc, "按中文词项和双字片段计算轻量语义相关度，并支持设备/故障/风险过滤。")
    add_number(doc, "返回引用编号、来源、章节、片段与相关度；无结果时明确返回空，不虚构文献。")
    add_figure(doc, "06-rag-evidence.png", "图3  辅助研判页的RAG命中、来源、相关度和引用片段")
    doc.add_heading("受控MaintenanceAgent", level=2)
    doc.add_paragraph("Agent按固定能力边界调用设备状态、遥测、预警、历史、知识、库存、诊断、工单、通知和知识候选十类工具。最大步骤、超时、重复工单检测与默认人工确认防止失控；界面只显示工具步骤摘要和外部证据。")
    add_figure(doc, "07-agent-run.png", "图4  Agent执行：十项工具摘要、风险结论与人工确认状态")
    doc.add_heading("豆包Provider与多模态", level=2)
    doc.add_paragraph("豆包调用只在服务端读取密钥，执行超时、限流、Schema校验和一次修复重试；失败后切换规则Provider。多模态入口完成文件校验、图片展示、遥测与RAG关联；未配置视觉模型时明确不识别具体缺陷。")
    add_figure(doc, "08-multimodal-analysis.png", "图5  多模态辅助：图片、当前遥测、人工检查建议与限制说明")

    doc.add_page_break()
    doc.add_heading("5. 1号引风机故事化案例", level=1)
    steps = [
        "高负荷运行中振动和轴承温度持续升高，健康度降至68，触发二级预警。",
        "RAG召回轴承磨损、润滑不足和联轴器不对中的可追溯案例。",
        "Agent读取趋势、历史、库存与知识，建议24小时内检查。",
        "3D自动定位驱动端轴承；高风险场景显示温度82℃、振动5.2 mm/s、故障概率89%。",
        "库存检查通过后创建工单，按待接单→已接单→检修中→待验证→已完成推进。",
        "验证后健康度恢复到92、风险变健康、预警关闭、轴承库存6→5，并形成知识候选。",
    ]
    for step in steps:
        add_number(doc, step)
    add_figure(doc, "05-digital-twin-fault.png", "图6  引风机轴承温升：故障部件、核心风险指标与维修主操作")
    add_figure(doc, "11-work-order.png", "图7  轨道式工单状态与真实业务字段")

    doc.add_page_break()
    doc.add_heading("6. 语义数字孪生与模型库", level=1)
    doc.add_paragraph("模型库用于部件定位、故障高亮、剖视和动画联动，不追求厂家级机械复原。所有资产随构建发布，不依赖在线HDR、CDN或外部模型。")
    add_table(doc, ["模型", "LOD0三角面", "LOD1", "LOD2", "接入"], [
        ["引风机", "119,993", "29,990", "7,985", "12台设备中的风机类"],
        ["工业电机", "1,112", "704", "464", "模型库"],
        ["轴承组件", "2,132", "804", "452", "模型库"],
        ["弹性联轴器", "1,204", "724", "484", "模型库"],
        ["给水泵", "4,548", "2,028", "1,116", "FWP-001/002"],
        ["循环水泵", "3,880", "1,696", "928", "CWP-001/002、CLP-001"],
        ["炉排减速机", "2,320", "1,300", "652", "GRB-001"],
        ["空压机", "4,284", "1,932", "1,044", "ACP-001"],
        ["渗滤液泵", "2,940", "1,212", "708", "LCP-001"],
    ], [1.25, 1.15, 0.9, 0.9, 2.3])
    add_callout(doc, "引风机关键验证", "增强模型16/16关键语义节点通过；原模型仍可通过model=original回退；模型切换使用URL作为组件key并重置节点、材质、透明和选择状态。")

    doc.add_heading("7. 飞书原生协同", level=1)
    add_table(doc, ["能力", "实现", "边界"], [
        ["网页免登", "客户端检测、临时授权、服务端换取身份", "普通浏览器使用演示用户"],
        ["多维表格", "工单优先recordId读写，能力级配置", "未配置模块独立Mock"],
        ["机器人/卡片", "高风险通知、接单/暂缓、详情跳转", "发送失败不回滚工单"],
        ["事件/WebSocket", "challenge、Token/Encrypt Key、幂等、IntentRouter", "平台权限和发布需人工"],
        ["Aily/妙搭", "完整字段、工具、页面和验收指南", "企业后台创建需人工"],
    ], [1.2, 3.0, 2.3])

    doc.add_page_break()
    doc.add_heading("8. 工程质量、安全与性能", level=1)
    add_bullet(doc, "严格TypeScript、Zod写入校验、统一错误响应；不把内部堆栈返回给客户端。")
    add_bullet(doc, "Secret、Token和模型Key只存在服务端；前端构建关闭source map；日志只输出缺失变量名或脱敏错误。")
    add_bullet(doc, "Mock操作保存到当前浏览器localStorage并从固定数据重放；Feishu模式不把localStorage当正式数据源。")
    add_bullet(doc, "3D路由懒加载，模型有三档LOD；1366、1440、1920分辨率无横向溢出，Canvas宽度约680、754、1024像素。")
    add_bullet(doc, "npm audit生产依赖未发现漏洞；最终lint、typecheck、test、build和生产烟测结果见最终验收报告。")

    doc.add_heading("9. 落地路径与价值边界", level=1)
    add_number(doc, "比赛原型：固定种子Mock、规则健康度、本地RAG、语义3D和工单闭环。")
    add_number(doc, "单场站影子运行：接入历史和只读实时数据，完成测点、单位、时钟和数据质量治理。")
    add_number(doc, "受控协同：平台审批后启用飞书/Aily/妙搭和远程Provider，关键动作保持人工确认。")
    add_number(doc, "规模化：建立模型监控、权限审计、灾备、跨场站知识治理和真实收益评估。")
    add_callout(doc, "不宣称已实现收益", "告警到接单时长、按时闭环率、重复故障召回率和缺料拦截数是后续试点指标；真实经济收益需要企业历史基线、经营口径和对照验证。", RED)

    doc.add_heading("10. 结论", level=1)
    doc.add_paragraph("烽燧智守证明了工业AI可以不止生成一句诊断：它能够基于工况和数据形成证据，调用知识、库存和工单工具，把结果推进到责任、处置、验证和沉淀。比赛版本的优势是可运行、可解释、可离线、可逐步接入；生产化的前提则是真实数据、现场验证、平台审批与专业人员决策。")

    doc.core_properties.title = "烽燧智守——项目附加报告（最终版）"
    doc.core_properties.subject = "垃圾焚烧发电关键设备故障预警与智能运维"
    doc.core_properties.author = "烽燧智守项目组"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()
